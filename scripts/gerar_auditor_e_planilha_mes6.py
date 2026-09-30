import json
import re
import sys
from pathlib import Path
from collections import Counter
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

def gerar_arquivos():
    with open('temp/conferencia_mes6_plano_produtivo.json', 'r', encoding='utf-8') as f:
        dados = json.load(f)

    registros = dados['registros']

    # ---------------------------------------------------------
    # 1. GERAR SCRIPT JAVASCRIPT AUDITOR PARA O SIGATER
    # ---------------------------------------------------------
    coletum_js_array = []
    for r in registros:
        coletum_js_array.append({
            "coletum_id": r['coletum_id'],
            "tecnico": r['tecnico'],
            "beneficiario": r['beneficiario'],
            "cpf": re.sub(r'\D', '', r['cpf']),
            "cpf_formatado": r['cpf'],
            "municipio": r['municipio'].split('-')[0] if '-' in r['municipio'] else r['municipio'],
            "comunidade": r['comunidade'],
            "data": r['data'],
            "visita_n": r.get('visita_n', '')
        })

    js_template = f"""/**
 * ==============================================================================
 * 🔍 AUDITOR SIGATER vs COLETUM - ACOMPANHAMENTO DO PLANO PRODUTIVO (JUNHO / 2026 - MÊS 6)
 * ==============================================================================
 * Formulário Coletum: 37226 (Acompanhamento Plano Produtivo - Bahia sem Fome)
 * Total no Coletum em Junho/2026: {len(registros)} registros (163 beneficiários únicos)
 *
 * Como usar:
 * 1. No SIGATER, abra a tela "LISTAGEM - EXECUÇÕES" da atividade de Acompanhamento / Visita Técnica.
 * 2. Abra o Console do navegador (F12 -> Console).
 * 3. Cole este código e pressione [ENTER].
 * ==============================================================================
 */
(async function() {{
    console.clear();
    console.log('%c🚀 INICIANDO AUDITORIA ACOMPANHAMENTO PLANO PRODUTIVO (JUNHO/2026 - {len(registros)} REGISTROS)...', 'color: #00bcd4; font-size: 16px; font-weight: bold;');

    const coletumRegistros = {json.dumps(coletum_js_array, ensure_ascii=False, indent=2)};

    function normalizar(txt) {{
        if (!txt) return '';
        return String(txt).normalize('NFD').replace(/[\\u0300-\\u036f]/g, '').toUpperCase().trim();
    }}

    function similaridade(s1, s2) {{
        const n1 = normalizar(s1);
        const n2 = normalizar(s2);
        if (n1 === n2) return 1.0;
        if (!n1 || !n2) return 0.0;
        if (n1.includes(n2) || n2.includes(n1)) return 0.92;
        const p1 = n1.split(' '), p2 = n2.split(' ');
        if (p1.length >= 2 && p2.length >= 2) {{
            if (p1[0] === p2[0] && p1[p1.length - 1] === p2[p2.length - 1]) return 0.88;
        }}
        const set1 = new Set(p1), set2 = new Set(p2);
        const inter = new Set([...set1].filter(x => set2.has(x)));
        return inter.size / new Set([...set1, ...set2]).size;
    }}

    // 1. Identifica links das execuções na tela do SIGATER
    const linksLupa = Array.from(document.querySelectorAll('a[href*="/read/"], a[href*="cronograma_execucao/read"], table tbody tr a'));
    const urlsExecucoes = Array.from(new Set(
        linksLupa.map(a => a.href).filter(h => h && (h.includes('/read/') || h.includes('cronograma_execucao/read')))
    ));

    console.log(`📋 Total de execuções lidas na tela do SIGATER: ${{urlsExecucoes.length}}`);

    if (urlsExecucoes.length === 0) {{
        alert('Por favor, abra a tela "LISTAGEM - EXECUÇÕES" da atividade no SIGATER.');
        return;
    }}

    // Painel Visual Flutuante
    const antigo = document.getElementById('painelAuditoriaAcomp');
    if (antigo) antigo.remove();

    const div = document.createElement('div');
    div.id = 'painelAuditoriaAcomp';
    div.style.cssText = 'position:fixed;bottom:20px;right:20px;max-width:580px;max-height:85vh;overflow-y:auto;background:#1e1e2f;color:#fff;padding:20px;border-radius:12px;box-shadow:0 10px 30px rgba(0,0,0,0.8);z-index:999999;font-family:sans-serif;font-size:13px;border:2px solid #00bcd4;';
    div.innerHTML = `
        <div style="border-bottom:1px solid #444;padding-bottom:8px;margin-bottom:12px;">
            <strong style="color:#00bcd4;font-size:15px;">⏳ Lendo ${{urlsExecucoes.length}} Execuções no SIGATER...</strong>
        </div>
        <p id="statusProgresso" style="margin:0;color:#ccc;">Carregando dados dos beneficiários...</p>
        <div style="width:100%;background:#333;height:10px;border-radius:5px;margin-top:10px;overflow:hidden;">
            <div id="barraProgresso" style="width:0%;height:100%;background:#00bcd4;transition:width 0.2s;"></div>
        </div>
    `;
    document.body.appendChild(div);

    // 2. Leitura paralela das páginas de execução
    const sigaterLancados = [];
    const BATCH_SIZE = 10;
    let concluidos = 0;

    for (let i = 0; i < urlsExecucoes.length; i += BATCH_SIZE) {{
        const batch = urlsExecucoes.slice(i, i + BATCH_SIZE);
        await Promise.all(batch.map(async (url) => {{
            try {{
                const resp = await fetch(url);
                if (resp.ok) {{
                    const html = await resp.text();
                    const doc = new DOMParser().parseFromString(html, 'text/html');
                    
                    let nomeBeneficiario = '';
                    let cpfBeneficiario = '';

                    const trsPart = Array.from(doc.querySelectorAll('table tr'));
                    for (const tr of trsPart) {{
                        const texto = tr.innerText || '';
                        if (texto.includes('Dados da DAP') || texto.includes('DAP_MDA') || /\\d{{3}}\\.\\d{{3}}\\.\\d{{3}}-\\d{{2}}/.test(texto)) {{
                            const tds = Array.from(tr.querySelectorAll('td'));
                            tds.forEach(td => {{
                                const t = td.innerText.trim();
                                if (/\\d{{3}}\\.\\d{{3}}\\.\\d{{3}}-\\d{{2}}/.test(t)) {{
                                    cpfBeneficiario = t;
                                }} else if (t.split(' ').length >= 2 && t.length >= 6 && !t.includes('DAP') && !t.includes('SIGATER') && !t.includes('MDA')) {{
                                    nomeBeneficiario = t;
                                }}
                            }});
                        }}
                    }}

                    if (!nomeBeneficiario) {{
                        const pdfLinks = Array.from(doc.querySelectorAll('a[href*=".pdf"], [title*=".pdf"]')).map(a => a.innerText || a.title || '');
                        for (const pdf of pdfLinks) {{
                            const match = pdf.match(/([A-Z_]+)_-_(?:ATESTE|COLETUM)/i);
                            if (match) {{
                                nomeBeneficiario = match[1].replace(/_/g, ' ');
                                break;
                            }}
                        }}
                    }}

                    const codExec = url.match(/\\/read\\/(\\d+)/) ? url.match(/\\/read\\/(\\d+)/)[1] : '';

                    sigaterLancados.push({{
                        codExec,
                        url,
                        nome: nomeBeneficiario || 'Não Identificado',
                        cpf: cpfBeneficiario || '-',
                        htmlBruto: normalizar(html)
                    }});
                }}
            }} catch (err) {{
            }} finally {{
                concluidos++;
                const perc = Math.round((concluidos / urlsExecucoes.length) * 100);
                const elStatus = document.getElementById('statusProgresso');
                const elBarra = document.getElementById('barraProgresso');
                if (elStatus) elStatus.innerText = `Processando: ${{concluidos}} de ${{urlsExecucoes.length}} (${{perc}}%)...`;
                if (elBarra) elBarra.style.width = `${{perc}}%`;
            }}
        }}));
    }}

    // 3. Cruzamento
    const lancadosConfirmados = [];
    const pendentesNaoLancados = [];

    coletumRegistros.forEach(col => {{
        let achou = false;
        for (const sig of sigaterLancados) {{
            const cpfLimpo = (col.cpf || '').replace(/\\D/g, '');
            if (similaridade(col.beneficiario, sig.nome) >= 0.70 || sig.htmlBruto.includes(normalizar(col.beneficiario)) || (cpfLimpo.length >= 8 && sig.cpf.replace(/\\D/g, '').includes(cpfLimpo))) {{
                achou = true;
                break;
            }}
        }}
        if (achou) {{
            lancadosConfirmados.push(col);
        }} else {{
            pendentesNaoLancados.push(col);
        }}
    }});

    // 4. Exibição
    console.log('%c====================================================================', 'color: #888');
    console.log(`%c📊 RESULTADO DA AUDITORIA (JUNHO/2026 - ACOMPANHAMENTO PLANO PRODUTIVO):
- Meta no Coletum (Junho/2026): %c${{coletumRegistros.length}}%c
- Lançados no SIGATER: %c${{sigaterLancados.length}}%c
- Confirmados: %c${{lancadosConfirmados.length}}%c
- ⚠️ PENDENTES: %c${{pendentesNaoLancados.length}}%c`,
        'font-weight: bold; font-size: 14px; color: #fff;',
        'color: #00e676; font-weight: bold;', 'color: #fff;',
        'color: #00bcd4; font-weight: bold;', 'color: #fff;',
        'color: #29b6f6; font-weight: bold;', 'color: #fff;',
        'color: #ff1744; font-weight: bold; font-size: 16px;', 'color: #fff;'
    );
    console.log('%c====================================================================', 'color: #888');

    if (pendentesNaoLancados.length > 0) {{
        console.log(`%c⚠️ BENEFICIÁRIOS PENDENTES NO SIGATER:`, 'color: #ff5252; font-size: 15px; font-weight: bold;');
        console.table(pendentesNaoLancados.map((p, i) => ({{
            '#': i + 1,
            'Beneficiário': p.beneficiario,
            'CPF': p.cpf_formatado,
            'Técnico': p.tecnico,
            'Comunidade': p.comunidade,
            'Município': p.municipio,
            'Data': p.data,
            'Visita Nº': p.visita_n
        }})));
    }} else {{
        console.log('%c🎉 Todos os beneficiários constam como lançados no SIGATER!', 'color: #00e676; font-size: 16px; font-weight: bold;');
    }}

    // 5. Painel Flutuante Final
    div.style.border = pendentesNaoLancados.length > 0 ? '2px solid #ff5252' : '2px solid #00e676';
    div.innerHTML = `
        <div style="display:flex;justify-content:space-between;align-items:center;border-bottom:1px solid #444;padding-bottom:8px;margin-bottom:12px;">
            <strong style="color:${{pendentesNaoLancados.length > 0 ? '#ff5252' : '#00e676'}};font-size:15px;">
                ${{pendentesNaoLancados.length > 0 ? '⚠️ Pendências de Acompanhamento (Junho/2026)' : '✅ 100% Lançado no SIGATER'}}
            </strong>
            <button onclick="document.getElementById('painelAuditoriaAcomp').remove()" style="background:none;border:none;color:#fff;cursor:pointer;font-size:16px;">✖</button>
        </div>
        <p style="margin:0 0 10px 0;color:#ccc;">
            Total Coletum: <strong>${{coletumRegistros.length}}</strong> | SIGATER: <strong style="color:#00e676;">${{sigaterLancados.length}}</strong> | Pendentes: <strong style="color:#ff5252;font-size:14px;">${{pendentesNaoLancados.length}}</strong>
        </p>
        <div style="max-height:360px;overflow-y:auto;">
            ${{pendentesNaoLancados.map((p, i) => `
                <div style="background:#2a2b40;padding:10px;border-radius:6px;margin-bottom:8px;border-left:4px solid #ff5252;">
                    <strong style="color:#fff;font-size:13px;">${{i+1}}. ${{p.beneficiario}}</strong><br>
                    <small style="color:#bbb;">CPF: ${{p.cpf_formatado}} | ${{p.comunidade}} (${{p.municipio}})</small><br>
                    <small style="color:#ffd54f;">👤 Técnico: ${{p.tecnico}} • 📅 Data: ${{p.data}} (Visita Nº ${{p.visita_n}})</small>
                </div>
            `).join('')}}
        </div>
        ${{pendentesNaoLancados.length > 0 ? `
        <button id="btnCopiarPendentesAcomp" style="width:100%;margin-top:12px;padding:10px;background:#00bcd4;color:#000;border:none;border-radius:6px;font-weight:bold;cursor:pointer;font-size:13px;">
            📋 Copiar Lista dos ${{pendentesNaoLancados.length}} Pendentes
        </button>` : ''}}
    `;

    if (document.getElementById('btnCopiarPendentesAcomp')) {{
        document.getElementById('btnCopiarPendentesAcomp').onclick = function() {{
            const txt = pendentesNaoLancados.map((p, i) => `${{i+1}}. ${{p.beneficiario}} | CPF: ${{p.cpf_formatado}} | Téc: ${{p.tecnico}} - ${{p.comunidade}} (${{p.municipio}}) - Data: ${{p.data}} (Visita ${{p.visita_n}})`).join('\\n');
            navigator.clipboard.writeText(txt).then(() => alert(`Lista dos ${{pendentesNaoLancados.length}} pendentes copiada!`));
        }};
    }}
}})();
"""

    caminho_js = Path('scripts/archive/auditors_js/auditor_acompanhamento_plano_produtivo_mes6.js')
    with open(caminho_js, 'w', encoding='utf-8') as f:
        f.write(js_template)
    print(f"✅ Script auditor JS criado em: {caminho_js}")

    # ---------------------------------------------------------
    # 2. GERAR PLANILHA EXCEL FORMATADA COM ESTILOS PROFISSIONAIS
    # ---------------------------------------------------------
    wb = openpyxl.Workbook()
    
    # Cores
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    sub_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    white_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # ABA 1: RESUMO EXECUTIVO
    ws_resumo = wb.active
    ws_resumo.title = "Resumo Executivo"
    ws_resumo.views.sheetView[0].showGridLines = True

    ws_resumo.merge_cells("A1:E1")
    ws_resumo["A1"] = "RELATÓRIO DE CONFERÊNCIA - ACOMPANHAMENTO PLANO PRODUTIVO (JUNHO/2026 - MÊS 6)"
    ws_resumo["A1"].font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    ws_resumo["A1"].fill = header_fill
    ws_resumo["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_resumo.row_dimensions[1].height = 35

    ws_resumo["A3"] = "Métrica Geral"
    ws_resumo["B3"] = "Valor"
    ws_resumo["A3"].font = bold_font
    ws_resumo["B3"].font = bold_font
    ws_resumo["A3"].fill = sub_fill
    ws_resumo["B3"].fill = sub_fill

    kpis = [
        ("Total de Registros no Coletum (Mês 06/2026)", len(registros)),
        ("Beneficiários Únicos Atendidos", len(set(r['cpf'] for r in registros if r['cpf']))),
        ("Duplicidades de Envio Detectadas", len(dados.get('duplicidades', {}))),
        ("Documentação Física/Digital 100% Completa (Ateste + Coletum)", dados['status_pastas_resumo']['completo']),
        ("Pendências de Pasta / Arquivo Escaneado no Computador", dados['status_pastas_resumo']['sem_pasta'] + dados['status_pastas_resumo']['falta_ateste']),
    ]

    for idx, (label, val) in enumerate(kpis, start=4):
        ws_resumo[f"A{idx}"] = label
        ws_resumo[f"B{idx}"] = val
        ws_resumo[f"A{idx}"].font = regular_font
        ws_resumo[f"B{idx}"].font = bold_font
        ws_resumo[f"A{idx}"].border = thin_border
        ws_resumo[f"B{idx}"].border = thin_border
        if "Completa" in label:
            ws_resumo[f"B{idx}"].fill = green_fill
        elif "Pendências" in label or "Duplicidades" in label and val > 0:
            ws_resumo[f"B{idx}"].fill = red_fill

    # Tabela por Técnico
    row_tec_start = 11
    ws_resumo[f"A{row_tec_start}"] = "Técnico Responsável"
    ws_resumo[f"B{row_tec_start}"] = "Total Cadastros"
    ws_resumo[f"A{row_tec_start}"].font = bold_font
    ws_resumo[f"B{row_tec_start}"].font = bold_font
    ws_resumo[f"A{row_tec_start}"].fill = sub_fill
    ws_resumo[f"B{row_tec_start}"].fill = sub_fill

    r_idx = row_tec_start + 1
    for tec, total in dados['resumo_tecnicos'].items():
        ws_resumo[f"A{r_idx}"] = tec
        ws_resumo[f"B{r_idx}"] = total
        ws_resumo[f"A{r_idx}"].font = regular_font
        ws_resumo[f"B{r_idx}"].font = bold_font
        ws_resumo[f"A{r_idx}"].border = thin_border
        ws_resumo[f"B{r_idx}"].border = thin_border
        r_idx += 1

    # Tabela por Município
    row_mun_start = r_idx + 2
    ws_resumo[f"A{row_mun_start}"] = "Município"
    ws_resumo[f"B{row_mun_start}"] = "Total Cadastros"
    ws_resumo[f"A{row_mun_start}"].font = bold_font
    ws_resumo[f"B{row_mun_start}"].font = bold_font
    ws_resumo[f"A{row_mun_start}"].fill = sub_fill
    ws_resumo[f"B{row_mun_start}"].fill = sub_fill

    r_idx2 = row_mun_start + 1
    for mun, total in dados['resumo_municipios'].items():
        ws_resumo[f"A{r_idx2}"] = mun
        ws_resumo[f"B{r_idx2}"] = total
        ws_resumo[f"A{r_idx2}"].font = regular_font
        ws_resumo[f"B{r_idx2}"].font = bold_font
        ws_resumo[f"A{r_idx2}"].border = thin_border
        ws_resumo[f"B{r_idx2}"].border = thin_border
        r_idx2 += 1

    ws_resumo.column_dimensions["A"].width = 45
    ws_resumo.column_dimensions["B"].width = 20

    # ABA 2: TODOS OS 164 REGISTROS
    ws_all = wb.create_sheet(title="Todos os 164 Registros")
    ws_all.views.sheetView[0].showGridLines = True

    headers = [
        "Nº", "ID Coletum", "Técnico Responsável", "Beneficiário(a)", "CPF",
        "Município", "Comunidade", "Data Realização", "Nº Visita", "Status Documental Local"
    ]
    ws_all.append(headers)
    ws_all.row_dimensions[1].height = 25
    for col_idx in range(1, len(headers) + 1):
        cell = ws_all.cell(row=1, column=col_idx)
        cell.font = white_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for i, r in enumerate(registros, start=1):
        st_local = r.get('status_local', 'SEM_PASTA')
        st_txt = "✅ COMPLETO" if st_local == 'COMPLETO' else ("⚠️ FALTA ATESTE" if st_local == 'FALTA_ATESTE' else "❌ SEM PASTA MÊS 6")
        row_vals = [
            i,
            r['coletum_id'],
            r['tecnico'],
            r['beneficiario'],
            r['cpf'],
            r['municipio'],
            r['comunidade'],
            r['data'],
            f"Visita {r.get('visita_n', '')}",
            st_txt
        ]
        ws_all.append(row_vals)
        curr_row = i + 1
        ws_all.row_dimensions[curr_row].height = 20
        for c in range(1, len(row_vals) + 1):
            cell = ws_all.cell(row=curr_row, column=c)
            cell.font = regular_font
            cell.border = thin_border
            if c == 10:
                cell.font = bold_font
                if "COMPLETO" in st_txt:
                    cell.fill = green_fill
                elif "FALTA ATESTE" in st_txt:
                    cell.fill = yellow_fill
                else:
                    cell.fill = red_fill

    for col in ws_all.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_all.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # ABA 3: DUPLICIDADES
    ws_dup = wb.create_sheet(title="Duplicidades Coletum")
    ws_dup.views.sheetView[0].showGridLines = True
    dup_headers = ["Beneficiário", "CPF", "ID Coletum", "Data Realização", "Técnico", "Comunidade", "Observação"]
    ws_dup.append(dup_headers)
    ws_dup.row_dimensions[1].height = 25
    for col_idx in range(1, len(dup_headers) + 1):
        cell = ws_dup.cell(row=1, column=col_idx)
        cell.font = white_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    dup_row = 2
    for cpf, items in dados.get('duplicidades', {}).items():
        for item in items:
            ws_dup.append([
                item['beneficiario'],
                item['cpf'],
                item['coletum_id'],
                item['data'],
                item['tecnico'],
                item['comunidade'],
                "Enviado 2x no Coletum em Junho/2026"
            ])
            for c in range(1, len(dup_headers) + 1):
                cell = ws_dup.cell(row=dup_row, column=c)
                cell.font = regular_font
                cell.border = thin_border
                cell.fill = yellow_fill
            dup_row += 1

    for col in ws_dup.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_dup.column_dimensions[col_letter].width = max(max_len + 3, 15)

    caminho_excel = Path(r"C:\Users\CLIENTE\Desktop\BAHIA_SEM_FOME\weverton\Conferencia_Acompanhamento_Plano_Produtivo_Mes6.xlsx")
    wb.save(caminho_excel)
    print(f"✅ Planilha Excel gerada com sucesso em: {caminho_excel}")

if __name__ == '__main__':
    gerar_arquivos()
