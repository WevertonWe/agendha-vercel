import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Carregar dados já processados
df_col = pd.read_csv("coletum_caracterizacao_05_2026.csv", sep=";")

# Detalhes das 6 Pendentes:
# 1. Elaine Jorge da Silva Rodrigues (Migson - Sansaite)
# 2. Josineide da Cruz (Migson - Sansaite)
# 3. Amilton José Paiva da Silva (Caroline - Lagoa do José Alves)
# 4. Claudeane Souza Oliveira (Caroline - Lagoa do José Alves)
# 5. Elídia Barbosa da Silva (Caroline - Alto Vermelho)
# 6. Adimilson Simões da Silva (Caroline - Lagoa do José Alves)

pendentes_6 = [
    {
        "tecnico": "Migson Brayne Pamponet da Silva",
        "beneficiario": "Elaine Jorge da Silva Rodrigues",
        "cpf": "038.905.025-06",
        "comunidade": "Sansaite",
        "municipio": "Macururé",
        "data": "20/05/2026",
        "coletum_id": "31657.67",
        "status_pasta": "FALTA COLETUM NA PASTA",
        "diagnostico_motivo": "Possui Ateste assinado na pasta. Falta vincular o PDF do Coletum e lançar no SIGATER."
    },
    {
        "tecnico": "Migson Brayne Pamponet da Silva",
        "beneficiario": "Josineide da Cruz",
        "cpf": "025.926.705-80",
        "comunidade": "Sansaite",
        "municipio": "Macururé",
        "data": "12/05/2026",
        "coletum_id": "31657.41",
        "status_pasta": "PASTA NOMEADA 2024 (COMPLETA)",
        "diagnostico_motivo": "Possui Ateste e Coletum, porém a pasta estava com erro de digitação (12.05.2024). Lançar no SIGATER."
    },
    {
        "tecnico": "Caroline Evangelista de Queiroz",
        "beneficiario": "Amilton José Paiva da Silva",
        "cpf": "096.430.635-26",
        "comunidade": "Lagoa do José Alves",
        "municipio": "Abaré",
        "data": "07/05/2026",
        "coletum_id": "31616.6",
        "status_pasta": "COMPLETO",
        "diagnostico_motivo": "Atividade 100% pronta na pasta (Ateste + Coletum). Falta apenas lançar no SIGATER."
    },
    {
        "tecnico": "Caroline Evangelista de Queiroz",
        "beneficiario": "Claudeane Souza Oliveira",
        "cpf": "075.827.175-11",
        "comunidade": "Lagoa do José Alves",
        "municipio": "Abaré",
        "data": "11/05/2026",
        "coletum_id": "31616.8",
        "status_pasta": "COMPLETO",
        "diagnostico_motivo": "Atividade 100% pronta na pasta (Ateste + Coletum). Falta apenas lançar no SIGATER."
    },
    {
        "tecnico": "Caroline Evangelista de Queiroz",
        "beneficiario": "Elídia Barbosa da Silva",
        "cpf": "000.227.935-50",
        "comunidade": "Alto Vermelho",
        "municipio": "Abaré",
        "data": "07/05/2026",
        "coletum_id": "31616.7",
        "status_pasta": "COMPLETO",
        "diagnostico_motivo": "Atividade 100% pronta na pasta (Ateste + Coletum). Falta apenas lançar no SIGATER."
    },
    {
        "tecnico": "Caroline Evangelista de Queiroz",
        "beneficiario": "Adimilson Simões da Silva",
        "cpf": "200.983.548-41",
        "comunidade": "Lagoa do José Alves",
        "municipio": "Abaré",
        "data": "11/05/2026",
        "coletum_id": "31616.9",
        "status_pasta": "FALTA ATESTE NA PASTA",
        "diagnostico_motivo": "Possui Coletum preenchido. Falta escanear o ateste assinado e lançar no SIGATER."
    }
]

df_pend = pd.DataFrame(pendentes_6)

# Criar Relatório Excel
wb = openpyxl.Workbook()
wb.remove(wb.active)

header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
fill_blue = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
fill_red = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
fill_green = PatternFill(start_color="385723", end_color="385723", fill_type="solid")
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

def formatar(ws, fill):
    ws.views.sheetView[0].showGridLines = True
    for col_idx, col in enumerate(ws.iter_cols(min_row=1, max_row=ws.max_row), start=1):
        h = ws.cell(1, col_idx)
        h.font = header_font
        h.fill = fill
        h.alignment = Alignment(horizontal="center", vertical="center")
        
        max_l = len(str(h.value or ''))
        for c in col[1:]:
            c.border = thin_border
            max_l = max(max_l, len(str(c.value or '')))
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = min(max_l + 4, 60)
    ws.row_dimensions[1].height = 25

# ABA 1: AS 6 PENDÊNCIAS CRÍTICAS (FALTAM NO SIGATER)
ws1 = wb.create_sheet(title="6 Pendentes no SIGATER")
ws1.append(["Nº", "Técnico", "Beneficiário", "CPF", "Comunidade", "Município", "Data Execução", "ID Coletum", "Situação Pasta", "Diagnóstico e Ação Necessária"])
for idx, r in df_pend.iterrows():
    ws1.append([idx+1, r['tecnico'], r['beneficiario'], r['cpf'], r['comunidade'], r['municipio'], r['data'], r['coletum_id'], r['status_pasta'], r['diagnostico_motivo']])
formatar(ws1, fill_red)

# ABA 2: AS 45 CONFIRMADAS NO SIGATER (MIGSON)
ws2 = wb.create_sheet(title="45 Confirmadas (Migson)")
ws2.append(["Nº", "Técnico", "Beneficiário", "CPF", "Comunidade", "Município", "Data Execução", "ID Coletum", "Status SIGATER"])

df_col_migson = df_col[df_col['tecnico'].str.contains('Migson', na=False)].copy()
# Remove Elaine e Josineide para listar as 45 confirmadas
df_45 = df_col_migson[~df_col_migson['beneficiario'].str.contains('Elaine|Josineide', na=False, case=False)]

for idx, r in df_45.reset_index().iterrows():
    d_format = str(r['data_ativ']).split('T')[0] if 'T' in str(r['data_ativ']) else str(r['data_ativ'])
    ws2.append([idx+1, r['tecnico'], r['beneficiario'], r['cpf'], r['comunidade'], r['municipio'], d_format, r['answer_id'], "CONFIRMADO / ENVIADO NO SIGATER"])
formatar(ws2, fill_green)

# ABA 3: TODAS AS 51 DO COLETUM
ws3 = wb.create_sheet(title="Todas as 51 do Coletum")
ws3.append(["Nº", "Técnico", "Beneficiário", "CPF", "Comunidade", "Município", "Data", "ID Coletum", "Situação no SIGATER"])
for idx, r in df_col.reset_index().iterrows():
    d_format = str(r['data_ativ']).split('T')[0] if 'T' in str(r['data_ativ']) else str(r['data_ativ'])
    is_pend = any(p['beneficiario'].lower() in str(r['beneficiario']).lower() for p in pendentes_6)
    sit = "⚠️ PENDENTE DE LANÇAMENTO (FALTA LANÇAR)" if is_pend else "✅ LANÇADO / CONFIRMADO"
    ws3.append([idx+1, r['tecnico'], r['beneficiario'], r['cpf'], r['comunidade'], r['municipio'], d_format, r['answer_id'], sit])
formatar(ws3, fill_blue)

rel_path = "Auditoria_Caracterizacao_I_Mes_23_Maio_2026.xlsx"
wb.save(rel_path)
print("Relatorio gerado com sucesso:", rel_path)
